"""Excel-driven twin of scripts/sheets_sync.py.

Same claim / progress / complete / fail semantics and the same column
names as the Google Sheets dashboard (apps_script/Code.gs COLUMNS, plus
one added column, "Facebook Status"), but operating on a local .xlsx
file instead of the Apps Script Web App bridge - no Google account,
no GAS_WEBAPP_URL / GAS_SHARED_SECRET needed. This is the input source
for the Muse + Hugging Face demo flow (see config/engine.json and
docs/DEMO_RUNBOOK.md). The Sheets flow stays untouched as a fallback.

Subcommands:
  show             Print all rows as JSON.
  claim            Claim the oldest Pending/blank row; prints
                   key=value lines (has_row, row_number, title, idea,
                   aspect_ratio, video_length, run_folder, today).
  progress         --row N --percent P    Update a row's Progress bar.
  update           --row N --set Col=Val [--set Col=Val ...]
  complete         --row N --video-url URL   Mark Completed, store URL,
                   set Approval Status back to Pending for human review.
  fail             --row N --reason "..."    Mark Failed with a reason.

File location: --file PATH, else $EXCEL_PATH, else <project>/dashboard.xlsx
"""
import argparse
import datetime
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from env_config import PIPELINE_ROOT

COLUMNS = [
    "Video Title", "Description", "Video Length", "Aspect Ratio", "Status",
    "Progress", "Approval Status", "Upload Status", "YouTube Status",
    "Facebook Status", "Last Updated", "Video URL", "Failure Reason",
]
SHEET_NAME = "Sheet1"
STALE_MINUTES = 25


def default_path():
    return os.environ.get("EXCEL_PATH") or os.path.join(PIPELINE_ROOT, "dashboard.xlsx")


def _now():
    return datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _bar(percent: int) -> str:
    filled = round(percent / 10)
    return "\u2593" * filled + "\u2591" * (10 - filled) + f" {percent}%"


def load(path):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    return wb, ws


def get_rows(ws):
    rows = []
    for r in range(2, ws.max_row + 1):
        entry = {"row_number": r}
        empty = True
        for j, col in enumerate(COLUMNS, start=1):
            v = ws.cell(row=r, column=j).value
            entry[col] = "" if v is None else str(v)
            if entry[col]:
                empty = False
        if not empty:
            rows.append(entry)
    return rows


def update_row(ws, row_number, values):
    for col, val in values.items():
        if col not in COLUMNS:
            raise ValueError(f"Unknown column: {col}")
        ws.cell(row=row_number, column=COLUMNS.index(col) + 1).value = val


def cmd_show(args):
    _, ws = load(args.file)
    print(json.dumps(get_rows(ws), indent=2, ensure_ascii=False))


def cmd_claim(args):
    wb, ws = load(args.file)
    rows = get_rows(ws)

    for row in rows:  # stale In-Progress recovery, same rule as sheets_sync
        if row["Status"] == "In Progress":
            try:
                ts = datetime.datetime.strptime(
                    row.get("Last Updated", ""), "%Y-%m-%dT%H:%M:%SZ"
                ).replace(tzinfo=datetime.timezone.utc)
                age_min = (datetime.datetime.now(datetime.timezone.utc) - ts).total_seconds() / 60
            except ValueError:
                age_min = STALE_MINUTES + 1
            if age_min > STALE_MINUTES:
                update_row(ws, row["row_number"], {
                    "Status": "Pending",
                    "Progress": "previous attempt stalled, retrying",
                    "Last Updated": _now(),
                })
                row["Status"] = "Pending"

    candidates = [r for r in rows if r["Status"] in ("", "Pending")]
    if not candidates:
        print("has_row=false")
        return
    row = candidates[0]
    row_number = row["row_number"]
    title = row["Video Title"]
    description = row.get("Description", "")
    idea = f"{title}. {description}" if description else title

    if "9:16" in row["Aspect Ratio"]:
        aspect_ratio = "9:16"
    elif "16:9" in row["Aspect Ratio"]:
        aspect_ratio = "16:9"
    else:
        update_row(ws, row_number, {
            "Status": "Completed",
            "Progress": "unsupported aspect ratio (only 9:16 or 16:9)",
            "Last Updated": _now(),
        })
        wb.save(args.file)
        print("has_row=false")
        return

    today = datetime.date.today().isoformat()
    run_folder = f"{today}-row{row_number}"
    update_row(ws, row_number, {
        "Status": "In Progress",
        "Progress": _bar(5),
        "Last Updated": _now(),
    })
    wb.save(args.file)

    print("has_row=true")
    print(f"row_number={row_number}")
    print(f"title={title}")
    print(f"idea={idea}")
    print(f"aspect_ratio={aspect_ratio}")
    print(f"video_length={row['Video Length'] or '30s'}")
    print(f"run_folder={run_folder}")
    print(f"today={today}")


def cmd_progress(args):
    wb, ws = load(args.file)
    update_row(ws, args.row, {"Progress": _bar(args.percent), "Last Updated": _now()})
    wb.save(args.file)
    print(f"row {args.row}: {_bar(args.percent)}")


def cmd_update(args):
    wb, ws = load(args.file)
    values = {}
    for item in args.set:
        key, _, val = item.partition("=")
        values[key.strip()] = val
    values["Last Updated"] = _now()
    update_row(ws, args.row, values)
    wb.save(args.file)
    print(f"row {args.row} updated: {values}")


def cmd_complete(args):
    wb, ws = load(args.file)
    update_row(ws, args.row, {
        "Status": "Completed",
        "Progress": _bar(100),
        "Approval Status": "Pending",
        "Video URL": args.video_url,
        "Failure Reason": "",
        "Last Updated": _now(),
    })
    wb.save(args.file)
    print(f"row {args.row} completed: {args.video_url}")


def cmd_fail(args):
    wb, ws = load(args.file)
    update_row(ws, args.row, {
        "Status": "Failed",
        "Failure Reason": args.reason,
        "Last Updated": _now(),
    })
    wb.save(args.file)
    print(f"row {args.row} failed: {args.reason}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", default=default_path())
    sub = parser.add_subparsers(dest="cmd", required=True)

    sub.add_parser("show")
    sub.add_parser("claim")

    p = sub.add_parser("progress")
    p.add_argument("--row", type=int, required=True)
    p.add_argument("--percent", type=int, required=True)

    p = sub.add_parser("update")
    p.add_argument("--row", type=int, required=True)
    p.add_argument("--set", action="append", required=True)

    p = sub.add_parser("complete")
    p.add_argument("--row", type=int, required=True)
    p.add_argument("--video-url", required=True)

    p = sub.add_parser("fail")
    p.add_argument("--row", type=int, required=True)
    p.add_argument("--reason", required=True)

    args = parser.parse_args()
    {
        "show": cmd_show, "claim": cmd_claim, "progress": cmd_progress,
        "update": cmd_update, "complete": cmd_complete, "fail": cmd_fail,
    }[args.cmd](args)
